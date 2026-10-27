"""Read the user's weekly power ratings from their model workbooks into build/ratings/{nfl,cfb}.json.
Output per team (dashboard team names): {"wk": {"0": rating after week 0 (= preseason + wk0 adj), "1": ..., }, "current": ...}
A week-N game is played by teams rated wk[N-1] going in. Only weeks with entered adjustments are kept.
usage: read_ratings.py nfl NFL_MODEL.xlsx | read_ratings.py cfb CFB_MODEL.xlsx"""
import sys, json, re, os, warnings
warnings.filterwarnings('ignore')
import openpyxl
LG, XL = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(XL, data_only=True, read_only=True)
rows = list(wb['Power Ratings History'].iter_rows(values_only=True))
hdr = rows[3]; col = {h: i for i, h in enumerate(hdr) if h}
data = [r for r in rows[4:] if r and r[0]]
# last week with any adjustment entered
last = max([w for w in range(0, 21) if f'Wk {w} Adj' in col and any(isinstance(r[col[f'Wk {w} Adj']], (int, float)) for r in data)] + [0])
if LG == 'nfl':
    sys.path.insert(0, os.path.join(HERE, 'nfl_full')); from nfl_season import NAME
    name = lambda x: NAME.get({'LAR': 'LA', 'JAC': 'JAX', 'WSH': 'WAS'}.get(x, x))   # note rows (e.g. injury memos) -> None
    teams = set(NAME.values())
else:
    teams = set(json.load(open(os.path.join(HERE, 'nfl_rebuild/cfb_team_list.json'))))
    nm = {str(a).strip(): str(b).strip() for a, b in list(wb['Name Map'].iter_rows(values_only=True))[1:] if a and b}
    FIX = {'Miami (FL)': 'Miami', 'Hawaii': "Hawai'i", 'JMU': 'James Madison', 'USF': 'South Florida', 'ECU': 'East Carolina',
           'FAU': 'Florida Atlantic', 'WKU': 'Western Kentucky', 'App. St.': 'App State', 'FIU': 'Florida International',
           'ULM': 'UL Monroe', 'San Jose St.': 'San José State', 'UMass': 'Massachusetts'}
    def name(x):
        x = str(x).strip()
        for cand in (x, nm.get(x), FIX.get(x), x.replace(' St.', ' State'), (nm.get(x) or '').replace(' St.', ' State')):
            if cand and cand in teams: return cand
        return None
out = {}; miss = []
for r in data:
    t = name(r[0])
    if t is None:
        if not isinstance(r[col['Current Rating']], (int, float)): continue   # not a team row
        miss.append(r[0]); continue
    wk = {str(w): float(r[col[f'Wk {w} Rating']]) for w in range(0, last + 1) if isinstance(r[col[f'Wk {w} Rating']], (int, float))}
    out[t] = {'wk': wk, 'current': float(r[col['Current Rating']])}
json.dump({'lastWeekGraded': last, 'teams': out}, open(os.path.join(HERE, f'ratings/{LG}.json'), 'w'), indent=0)
print(f'{LG}: {len(out)} teams matched, last graded week {last}; unmatched workbook names: {miss}; dashboard teams without a rating: {sorted(teams - set(out))[:20]}')
